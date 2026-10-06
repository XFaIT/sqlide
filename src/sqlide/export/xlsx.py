"""Excel .xlsx via openpyxl write-only mode (constant memory).

Handles Excel limits: 1,048,576 rows per sheet (continues on 'Name 2', ...), 32,767 chars
per cell, control characters Excel rejects, and the 15-digit float precision (larger
integers / decimals are written as text so no digits are silently lost).
"""

from __future__ import annotations

import datetime as dt
import math
import re
from collections.abc import Iterator
from decimal import Decimal
from itertools import chain, islice
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from sqlide.db.result import Column
from sqlide.export.base import Exporter, ExportOptions, Row, register

MAX_ROWS_PER_SHEET = 1_048_576 - 1  # minus the header row
MAX_CELL_CHARS = 32_767
SAMPLE_ROWS = 100
MAX_PRECISE = 10**15
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")
_FORMATS = {dt.datetime: "yyyy-mm-dd hh:mm:ss", dt.date: "yyyy-mm-dd", dt.time: "hh:mm:ss"}


def sheet_title(name: str, index: int) -> str:
    base = _BAD_SHEET.sub("_", name).strip("'") or "Result"
    suffix = f" {index}" if index > 1 else ""
    return base[: 31 - len(suffix)] + suffix


def to_cell_value(ws: Any, v: Any) -> Any:
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v if abs(v) < MAX_PRECISE else str(v)
    if isinstance(v, float):
        return v if math.isfinite(v) else str(v)
    if isinstance(v, Decimal):
        return float(v) if v.is_finite() and len(v.as_tuple().digits) <= 15 else str(v)
    for typ, fmt in _FORMATS.items():  # datetime before date: it is a subclass
        if isinstance(v, typ):
            cell = WriteOnlyCell(ws, value=v)
            cell.number_format = fmt
            return cell
    if isinstance(v, bytes):
        return "0x" + v.hex()
    return _ILLEGAL.sub("", str(v))[:MAX_CELL_CHARS]


@register
class XlsxExporter(Exporter):
    name, label, extension = "xlsx", "Excel (.xlsx)", ".xlsx"

    def _write(
        self, columns: list[Column], rows: Iterator[Row], tmp: Path, opts: ExportOptions
    ) -> int:
        wb = Workbook(write_only=True)
        total = in_sheet = sheet_no = 0
        ws: Any = None
        names = [c.name for c in columns]
        sample = list(islice(rows, SAMPLE_ROWS))
        widths = [
            min(60, max(len(n) + 2, *(len(str(r[i])) if r[i] is not None else 0 for r in sample)))
            for i, n in enumerate(names)
        ]

        def new_sheet() -> Any:
            nonlocal sheet_no
            sheet_no += 1
            sheet = wb.create_sheet(sheet_title(opts.sheet_name, sheet_no))
            for i, w in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(i)].width = w + 1
            if opts.header:
                sheet.freeze_panes = "A2"
                bold = Font(bold=True)
                cells = []
                for n in names:
                    cell = WriteOnlyCell(sheet, value=_ILLEGAL.sub("", n))
                    cell.font = bold
                    cells.append(cell)
                sheet.append(cells)
            return sheet

        def finish(sheet: Any, n_rows: int) -> None:
            if opts.header and names and n_rows:
                sheet.auto_filter.ref = f"A1:{get_column_letter(len(names))}{n_rows + 1}"

        ws = new_sheet()
        for row in chain(sample, rows):
            if in_sheet >= MAX_ROWS_PER_SHEET:
                finish(ws, in_sheet)
                ws, in_sheet = new_sheet(), 0
            ws.append([to_cell_value(ws, v) for v in row])
            in_sheet += 1
            total += 1
        finish(ws, in_sheet)
        wb.save(tmp)
        return total
