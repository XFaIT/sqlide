import csv
import datetime as dt
import json
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from sqlide import export
from sqlide.db.result import Column
from sqlide.export import ExportCancelled, ExportOptions, xlsx

COLS = [Column("id", "INT", 4), Column("Name", "VARCHAR", 12), Column("when", "TIMESTAMP", 93)]
ROWS = [
    (1, "plain", dt.datetime(2026, 1, 2, 3, 4, 5)),
    (2, 'q"uote, comma\nnewline', None),
    (3, None, dt.datetime(2026, 12, 31, 23, 59, 59)),
]


def run(name, tmp_path, rows=ROWS, cols=COLS, **opt):
    ex = export.get(name)
    path = tmp_path / f"out{ex.extension}"
    n = ex.write(cols, iter(rows), path, ExportOptions(**opt))
    return n, path


def test_registry_has_all_formats():
    assert {e.name for e in export.all_exporters()} == {
        "csv", "tsv", "json", "jsonl", "markdown", "sql", "html", "xlsx",
    }  # fmt: skip
    with pytest.raises(KeyError):
        export.get("pdf")


def test_csv_roundtrip_with_header_and_tricky_cells(tmp_path):
    n, path = run("csv", tmp_path)
    assert n == 3
    got = list(csv.reader(path.open(newline="", encoding="utf-8")))
    assert got[0] == ["id", "Name", "when"]
    assert got[2] == ["2", 'q"uote, comma\nnewline', ""]
    assert got[3] == ["3", "", "2026-12-31 23:59:59"]


def test_csv_options_no_header_delimiter_bom(tmp_path):
    _, path = run("csv", tmp_path, header=False, delimiter=";", bom=True)
    raw = path.read_bytes()
    assert (
        raw.startswith(b"\xef\xbb\xbf") and b"1;plain;" in raw and b"id" not in raw.split(b"\n")[0]
    )


def test_tsv_forces_tab(tmp_path):
    _, path = run("tsv", tmp_path, delimiter=";")
    assert path.read_text().splitlines()[0] == "id\tName\twhen"


def test_json_array_and_empty(tmp_path):
    _, path = run("json", tmp_path)
    data = json.loads(path.read_text())
    assert data[0] == {"id": 1, "Name": "plain", "when": "2026-01-02 03:04:05"}
    assert data[2]["Name"] is None
    _, empty = run("json", tmp_path, rows=[])
    assert json.loads(empty.read_text()) == []


def test_jsonl(tmp_path):
    _, path = run("jsonl", tmp_path)
    lines = path.read_text().splitlines()
    assert len(lines) == 3 and json.loads(lines[1])["id"] == 2


def test_markdown(tmp_path):
    _, path = run("markdown", tmp_path)
    lines = path.read_text().splitlines()
    assert lines[0] == "| id | Name | when |" and lines[1] == "| --- | --- | --- |"
    assert 'q"uote, comma<br>newline' in lines[3]


def test_html_escapes_and_marks_nulls(tmp_path):
    _, path = run("html", tmp_path, rows=[(1, "<b>&</b>", None)])
    text = path.read_text()
    assert "&lt;b&gt;&amp;&lt;/b&gt;" in text and '<td class="null"></td>' in text


def test_sql_insert_single_and_batched(tmp_path):
    cols = [Column("id", "INT", 4), Column("Mixed Name", "VARCHAR", 12)]
    rows = [(1, "o'k"), (2, None), (3, "z")]
    _, path = run("sql", tmp_path, rows=rows, cols=cols, table_name="public.t")
    lines = path.read_text().splitlines()
    assert lines[0] == 'INSERT INTO public.t (id, "Mixed Name") VALUES'
    assert "(1, 'o''k')" in path.read_text() and len(lines) == 6  # head + row, three times
    _, batched = run("sql", tmp_path, rows=rows, cols=cols, batch_size=2)
    text = batched.read_text()
    assert text.count("INSERT INTO") == 2 and "(1, 'o''k'),\n  (2, NULL);" in text


def test_xlsx_roundtrip_types_header_and_filter(tmp_path):
    cols = [Column(n, "", 0) for n in ("i", "f", "d", "s", "dt", "b", "big", "dec", "n")]
    row = (7, 1.5, Decimal("2.25"), "txt\x01x", dt.datetime(2026, 1, 2, 3, 4, 5), True,
           10**20, Decimal("1.23456789012345678"), None)  # fmt: skip
    n, path = run("xlsx", tmp_path, rows=[row], cols=cols, sheet_name="My/Sheet")
    assert n == 1
    ws = load_workbook(path)["My_Sheet"]
    header = [c.value for c in ws[1]]
    assert header == [c.name for c in cols] and ws[1][0].font.bold
    vals = [c.value for c in ws[2]]
    assert vals[:4] == [7, 1.5, 2.25, "txtx"]  # control char stripped, Decimal as number
    assert (
        vals[4] == dt.datetime(2026, 1, 2, 3, 4, 5)
        and ws[2][4].number_format == "yyyy-mm-dd hh:mm:ss"
    )
    assert vals[5] is True
    assert vals[6] == str(10**20) and vals[7] == "1.23456789012345678"  # no silent digit loss
    assert vals[8] is None
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:I2"


def test_xlsx_splits_sheets_at_the_row_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx, "MAX_ROWS_PER_SHEET", 5)
    cols = [Column("n", "INT", 4)]
    n, path = run("xlsx", tmp_path, rows=[(i,) for i in range(12)], cols=cols)
    wb = load_workbook(path)
    assert n == 12 and wb.sheetnames == ["Result", "Result 2", "Result 3"]
    assert [ws.max_row for ws in wb] == [6, 6, 3]  # header + rows each
    assert wb["Result 3"][1][0].value == "n"


def test_xlsx_long_text_is_truncated_not_rejected(tmp_path):
    _, path = run("xlsx", tmp_path, rows=[("x" * 40000,)], cols=[Column("s", "", 12)])
    assert len(load_workbook(path).active[2][0].value) == xlsx.MAX_CELL_CHARS


def test_cancel_leaves_no_file_and_no_partial(tmp_path):
    ex = export.get("csv")
    path = tmp_path / "x.csv"
    calls = iter([False, False, True])
    with pytest.raises(ExportCancelled):
        ex.write(COLS, iter(ROWS), path, ExportOptions(should_cancel=lambda: next(calls)))
    assert list(tmp_path.iterdir()) == []


def test_failure_midway_keeps_existing_file_intact(tmp_path):
    ex = export.get("csv")
    path = tmp_path / "x.csv"
    path.write_text("old")

    def rows():
        yield ROWS[0]
        raise RuntimeError("db died")

    with pytest.raises(RuntimeError):
        ex.write(COLS, rows(), path, ExportOptions())
    assert path.read_text() == "old" and [p.name for p in tmp_path.iterdir()] == ["x.csv"]


def test_progress_callback(tmp_path):
    seen = []
    ex = export.get("csv")
    ex.write(
        [Column("n", "INT", 4)],
        iter([(i,) for i in range(2500)]),
        tmp_path / "p.csv",
        ExportOptions(progress=seen.append),
    )
    assert seen == [1000, 2000]


def test_creates_missing_directories(tmp_path):
    ex = export.get("csv")
    target = tmp_path / "a" / "b" / "x.csv"
    ex.write(COLS, iter(ROWS), target, ExportOptions())
    assert target.exists()
