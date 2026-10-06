"""Whole-result export: stream a query through a real (H2) session straight into files."""

import csv

import pytest
from openpyxl import load_workbook

from sqlide import export
from sqlide.db.result import DbError
from sqlide.export import ExportOptions
from sqlide.export.service import export_query, export_rows


async def test_stream_exports_every_row_beyond_the_first_page(session, tmp_path):
    sql = "select x, 'row ' || x as label from system_range(1, 1200)"
    n = await export_query(export.get("csv"), session, sql, tmp_path / "all.csv", ExportOptions())
    assert n == 1200
    rows = list(csv.reader((tmp_path / "all.csv").open(newline="")))
    assert rows[0] == ["X", "LABEL"] and rows[-1] == ["1200", "row 1200"] and len(rows) == 1201


async def test_stream_to_xlsx(session, tmp_path):
    sql = "select x, x * 1.5 as v from system_range(1, 3000)"
    n = await export_query(export.get("xlsx"), session, sql, tmp_path / "a.xlsx", ExportOptions())
    ws = load_workbook(tmp_path / "a.xlsx").active
    assert n == 3000 and ws.max_row == 3001 and ws[3001][0].value == 3000


async def test_stream_rejects_statements_without_result_set(session, tmp_path):
    await session.execute("create table t(a int)")
    with pytest.raises(DbError, match="result set"):
        await export_query(
            export.get("csv"),
            session,
            "insert into t values (1)",
            tmp_path / "x.csv",
            ExportOptions(),
        )
    assert not (tmp_path / "x.csv").exists()


async def test_session_usable_after_stream(session, tmp_path):
    await export_query(
        export.get("csv"),
        session,
        "select x from system_range(1, 5000)",
        tmp_path / "a.csv",
        ExportOptions(),
    )
    assert (await session.execute("select 1")).items[0].rows == [(1,)]


async def test_export_rows_in_memory(tmp_path):
    from sqlide.db.result import Column

    n = await export_rows(
        export.get("json"),
        [Column("a", "INT", 4)],
        [(1,), (2,)],
        tmp_path / "m.json",
        ExportOptions(),
    )
    assert n == 2
