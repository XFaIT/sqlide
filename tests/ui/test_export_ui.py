"""Export dialog and execution, driven through the real app against H2."""

import csv

from openpyxl import load_workbook
from textual.widgets import Input, RadioButton, RadioSet, Select, Static, Switch

from sqlide.app import SqlideApp
from sqlide.config.connections import Connection
from sqlide.ui.screens.export_dialog import ExportScreen
from sqlide.ui.screens.main import MainScreen
from tests.ui.helpers import connect_first, console, grids, wait_for

URL = "jdbc:h2:mem:{};DB_CLOSE_DELAY=-1"


async def run_query(pilot, app, sql):
    console(app).editor.text = sql
    await pilot.press("f5")
    await wait_for(pilot, lambda: len(grids(app)) == 1 and not console(app).running)
    grid = grids(app)[0]
    grid.focus()
    return grid


async def open_export(pilot, app):
    await pilot.press("e")
    await wait_for(pilot, lambda: isinstance(app.screen, ExportScreen))
    return app.screen


async def test_export_whole_result_reruns_query_and_streams(make_ws, tmp_path):
    ws = make_ws(Connection("h", "h2", URL.format("ex1")))
    ws.settings.fetch_size = 50
    app = SqlideApp(ws)
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        grid = await run_query(pilot, app, "select x, 'v' || x as label from system_range(1, 200)")
        assert len(grid.model) == 50 and grid.source is not None
        dlg = await open_export(pilot, app)
        assert dlg.query_one("#scope", RadioSet).pressed_button.id == "all"  # more rows exist
        out = tmp_path / "all.csv"
        dlg.query_one("#path", Input).value = str(out)
        await pilot.click("#export")
        await wait_for(pilot, lambda: out.exists() and not console(app)._exporting)
        rows = list(csv.reader(out.open(newline="")))
        assert rows[0] == ["X", "LABEL"] and len(rows) == 201 and rows[-1] == ["200", "v200"]
        assert isinstance(app.screen, MainScreen)
        assert len(grid.model) == 50 and grid.source is not None  # the grid's cursor untouched


async def test_export_view_respects_sort_and_filter_and_swaps_extension(make_ws, tmp_path):
    ws = make_ws(Connection("h", "h2", URL.format("ex2")))
    app = SqlideApp(ws)
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        grid = await run_query(pilot, app, "select x, 'n' || x as nm from system_range(1, 30)")
        grid.set_filter("n2")  # 2, 20..29
        await pilot.press("s", "s")  # x descending
        dlg = await open_export(pilot, app)
        assert dlg.query_one("#scope", RadioSet).pressed_button.id == "view"
        path = dlg.query_one("#path", Input)
        assert path.value.endswith(".csv")
        dlg.query_one("#format", Select).value = "xlsx"
        await pilot.pause()
        assert path.value.endswith(".xlsx")  # extension followed the format
        out = tmp_path / "view.xlsx"
        path.value = str(out)
        await pilot.click("#export")
        await wait_for(pilot, lambda: out.exists() and not console(app)._exporting)
        ws_ = load_workbook(out).active
        assert [r[0].value for r in ws_.iter_rows(min_row=2)] == [
            29,
            28,
            27,
            26,
            25,
            24,
            23,
            22,
            21,
            20,
            2,
        ]


async def test_export_selection_scope(make_ws, tmp_path):
    app = SqlideApp(make_ws(Connection("h", "h2", URL.format("ex3"))))
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        await run_query(pilot, app, "select x, x * 10 as y, 'z' as z from system_range(1, 5)")
        await pilot.press("shift+down", "shift+down", "shift+right")  # 3 rows x 2 cols
        dlg = await open_export(pilot, app)
        dlg.query_one("#selection", RadioButton).value = True
        out = tmp_path / "sel.json"
        dlg.query_one("#format", Select).value = "json"
        await pilot.pause()
        dlg.query_one("#path", Input).value = str(out)
        await pilot.click("#export")
        await wait_for(pilot, lambda: out.exists() and not console(app)._exporting)
        import json

        assert json.loads(out.read_text()) == [
            {"X": 1, "Y": 10}, {"X": 2, "Y": 20}, {"X": 3, "Y": 30},
        ]  # fmt: skip


async def test_overwrite_is_guarded(make_ws, tmp_path):
    app = SqlideApp(make_ws(Connection("h", "h2", URL.format("ex4"))))
    existing = tmp_path / "keep.csv"
    existing.write_text("precious")
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        await run_query(pilot, app, "select 1 as a")
        dlg = await open_export(pilot, app)
        dlg.query_one("#path", Input).value = str(existing)
        await pilot.click("#export")
        await pilot.pause()
        assert isinstance(app.screen, ExportScreen)  # still open
        assert "exists" in str(dlg.query_one("#error", Static).render())
        assert existing.read_text() == "precious"
        dlg.query_one("#overwrite", Switch).value = True
        await pilot.click("#export")
        await wait_for(
            pilot, lambda: not console(app)._exporting and existing.read_text() != "precious"
        )
        assert existing.read_text().splitlines() == ["A", "1"]


async def test_export_failure_is_reported_not_fatal(make_ws, tmp_path):
    app = SqlideApp(make_ws(Connection("h", "h2", URL.format("ex5"))))
    blocker = tmp_path / "file"
    blocker.write_text("x")  # a file where a directory is needed
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        await run_query(pilot, app, "select 1 as a")
        dlg = await open_export(pilot, app)
        dlg.query_one("#path", Input).value = str(blocker / "sub" / "x.csv")
        await pilot.click("#export")
        await wait_for(pilot, lambda: console(app).status.message == "export failed")
        assert any("Export failed" in n.title for n in app._notifications)
        assert isinstance(app.screen, MainScreen)


async def test_export_dialog_cancel(make_ws):
    app = SqlideApp(make_ws(Connection("h", "h2", URL.format("ex6"))))
    async with app.run_test(size=(120, 50)) as pilot:
        await connect_first(pilot, app)
        await run_query(pilot, app, "select 1 as a")
        await open_export(pilot, app)
        await pilot.press("escape")
        await pilot.pause()
        assert isinstance(app.screen, MainScreen)
